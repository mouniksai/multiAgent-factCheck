"""
Professional Academic Presentation Generator (Minimal & Elegant)
==================================================================
Builds a minimalist, publication-grade 16:9 widescreen presentation
following academic consulting standards (McKinsey / Academic Conference style):
- Clean pure white background (#FFFFFF)
- Deep charcoal / slate typography (#111827 / #1F2937)
- Single subtle academic navy accent (#1E3A8A)
- Subtle grey gridlines (#E5E7EB) and subtle headers (#F3F4F6)
- Zero childish emoji spam, zero saturated rainbow boxes
"""

import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Professional Academic Color Palette (Minimal & Restrained)
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_BG_LIGHT = RGBColor(250, 250, 250)
COLOR_HEADER_BG = RGBColor(243, 244, 246)      # Subtle warm grey
COLOR_TEXT_PRIMARY = RGBColor(17, 24, 39)       # Dark charcoal / almost black
COLOR_TEXT_SECONDARY = RGBColor(75, 85, 99)     # Medium neutral slate
COLOR_TEXT_MUTED = RGBColor(107, 114, 128)      # Muted grey
COLOR_BORDER = RGBColor(229, 231, 235)          # Thin hairline border
COLOR_ACCENT = RGBColor(30, 58, 138)            # Deep academic navy
COLOR_ROW_ALT = RGBColor(249, 250, 251)         # Subtle zebra row


def set_pure_slide_background(slide):
    """Sets a clean, pure white background."""
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_WHITE
    bg.line.fill.background()
    return bg


def add_minimal_slide_header(slide, title_text, category_text="FOUNDATIONS OF ARTIFICIAL INTELLIGENCE"):
    """Adds a clean, professional slide header with subtle typography."""
    tb = slide.shapes.add_textbox(Inches(0.9), Inches(0.5), Inches(11.5), Inches(0.95))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p_cat = tf.paragraphs[0]
    p_cat.text = category_text.upper()
    p_cat.font.name = "Calibri"
    p_cat.font.size = Pt(9.5)
    p_cat.font.bold = True
    p_cat.font.color.rgb = COLOR_TEXT_MUTED

    p_title = tf.add_paragraph()
    p_title.text = title_text
    p_title.font.name = "Calibri"
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_TEXT_PRIMARY

    # Hairline divider
    divider = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.4), Inches(11.5), Inches(0.015))
    divider.fill.solid()
    divider.fill.fore_color.rgb = COLOR_BORDER
    divider.line.fill.background()


add_slide_header = add_minimal_slide_header


def format_minimal_table(table, col_widths, headers, rows):
    """Applies clean, publication-grade academic table formatting."""
    for idx, width in enumerate(col_widths):
        table.columns[idx].width = Inches(width)

    # Header Row
    for col_idx, header in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_HEADER_BG
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.LEFT
            p.font.name = "Calibri"
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = COLOR_TEXT_PRIMARY

    # Data Rows
    for row_idx, row_data in enumerate(rows):
        is_alt = (row_idx % 2 == 1)
        for col_idx, cell_value in enumerate(row_data):
            cell = table.cell(row_idx + 1, col_idx)
            cell.text = str(cell_value)
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_ROW_ALT if is_alt else COLOR_WHITE
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT
                p.font.name = "Calibri"
                p.font.size = Pt(10.5)
                p.font.color.rgb = COLOR_TEXT_SECONDARY


def create_minimal_deck(output_filename="Aletheia_Dialectical_Fact_Checking.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 1: Title Slide (Minimal Academic)
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_pure_slide_background(slide1)

    tb1 = slide1.shapes.add_textbox(Inches(1.5), Inches(1.8), Inches(10.3), Inches(4.2))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "FOUNDATIONS OF ARTIFICIAL INTELLIGENCE"
    p0.font.name = "Calibri"
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_ACCENT

    p1 = tf1.add_paragraph()
    p1.text = "Aletheia: Dialectical Multi-Agent Fact-Checking"
    p1.font.name = "Calibri"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_TEXT_PRIMARY

    p2 = tf1.add_paragraph()
    p2.text = "Adversarial RAG, Epistemic Calibration, and Citation Integrity"
    p2.font.name = "Calibri"
    p2.font.size = Pt(14)
    p2.font.color.rgb = COLOR_TEXT_SECONDARY

    # Thin horizontal accent rule
    rule = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.5), Inches(3.8), Inches(3.5), Inches(0.02))
    rule.fill.solid()
    rule.fill.fore_color.rgb = COLOR_ACCENT
    rule.line.fill.background()

    tb_auth = slide1.shapes.add_textbox(Inches(1.5), Inches(4.1), Inches(8.0), Inches(2.2))
    tf_auth = tb_auth.text_frame
    tf_auth.word_wrap = True

    p_by = tf_auth.paragraphs[0]
    p_by.text = "Presented by:"
    p_by.font.name = "Calibri"
    p_by.font.size = Pt(11)
    p_by.font.bold = True
    p_by.font.color.rgb = COLOR_TEXT_PRIMARY

    p_names = tf_auth.add_paragraph()
    p_names.text = "Team AI Fact-Checkers\n[Student Name 1] — [Roll No]\n[Student Name 2] — [Roll No]\n[Student Name 3] — [Roll No]"
    p_names.font.name = "Calibri"
    p_names.font.size = Pt(11.5)
    p_names.font.color.rgb = COLOR_TEXT_SECONDARY

    # =========================================================================
    # SLIDE 2: Multi-Agent RAG for Reliable Claim Verification
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_pure_slide_background(slide2)
    add_slide_header(slide2, "Multi-Agent RAG for Reliable Claim Verification", "SYSTEM OVERVIEW")

    tb2 = slide2.shapes.add_textbox(Inches(0.9), Inches(1.8), Inches(11.5), Inches(5.0))
    tf2 = tb2.text_frame
    tf2.word_wrap = True

    bullets = [
        "A multi-agent Retrieval-Augmented Generation (RAG) system designed to verify factual claims using live empirical web evidence.",
        "The system first decomposes complex assertions into smaller atomic sub-claims, eliminating multi-hop retrieval errors and query ambiguity.",
        "PRO and CON agents independently execute stance-conditioned web retrieval to discover evidence and engage in an adversarial debate across iterative rounds.",
        "An impartial Cross-Examiner agent audits debater turns for formal reasoning fallacies and isolates empirical consensus from active points of contention.",
        "A Citation Guardrail verifies that every cited evidence ID corresponds strictly to a verified record in the shared evidence pool before arbitration.",
        "An independent Supreme Judge analyzes verified evidence and domain credibility to output an epistemically calibrated verdict with explicit uncertainty notes."
    ]

    for idx, b_text in enumerate(bullets):
        p = tf2.paragraphs[0] if idx == 0 else tf2.add_paragraph()
        p.text = f"•  {b_text}"
        p.font.name = "Calibri"
        p.font.size = Pt(13.5)
        p.font.color.rgb = COLOR_TEXT_PRIMARY
        p.space_after = Pt(14)

    # =========================================================================
    # SLIDE 3: PEAS Formulation:
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_pure_slide_background(slide3)
    add_slide_header(slide3, "PEAS Formulation:", "REVIEW 1: 3 MARKS")

    peas_headers = ["Agent", "Performance Measure (P)", "Environment (E)", "Actuators (A)", "Sensors (S)"]
    peas_widths = [1.8, 2.7, 2.3, 2.4, 2.3]
    peas_rows = [
        ["Claim Normalizer", "Sub-claim atomicity, precision, logical independence", "Raw/unstructured claims text space", "Decomposed sub-claims list (JSON)", "User claim string"],
        ["PRO Debater", "Evidence relevance, citation precision, domain authority", "Live Web, debate transcript history", "Supporting queries & arguments [Ex]", "Search results, evidence pool, CON arguments"],
        ["CON Debater", "Counter-evidence rigor, refutation validity, ground truth", "Live Web, debate transcript history", "Refuting queries & arguments [Ex]", "Search results, evidence pool, PRO arguments"],
        ["Cross-Examiner", "Fallacy detection rate, dialectical synthesis neutrality", "Full multi-turn debate history", "Consensus/contention map, audit report", "Debater turns, citation mapping, pool"],
        ["Citation Guardrail", "100% citation validation, 0% hallucination tolerance", "Evidence registry & debate text", "Verification report, discount warnings", "Citation tags [E1], [E2], evidence keys"],
        ["Supreme Judge", "Verdict accuracy, epistemic calibration, nuance preservation", "Full debate + verified evidence pool", "TRUE / FALSE / UNVERIFIABLE + rationale", "Debate transcript, evidence, audit reports"]
    ]

    t_shape3 = slide3.shapes.add_table(len(peas_rows) + 1, len(peas_headers), Inches(0.9), Inches(1.7), Inches(11.5), Inches(4.9))
    format_minimal_table(t_shape3.table, peas_widths, peas_headers, peas_rows)

    # =========================================================================
    # SLIDE 4: Environment & Agent Analysis:
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_pure_slide_background(slide4)
    add_slide_header(slide4, "Environment & Agent Analysis:", "REVIEW 1: 3 MARKS")

    env_headers = ["Property", "Our Implementation & Formal Rationale"]
    env_widths = [3.2, 8.3]
    env_rows = [
        ["Stochastic", "Web retrieval latencies and LLM responses can vary between executions; temperature-conditioned sampling is non-deterministic."],
        ["Partially Observable", "Each debater initially observes only its own retrieved evidence prior to synchronization into the shared evidence pool."],
        ["Sequential", "Each debate round depends conditionally on previous PRO and CON arguments, building a stateful dialectic over time."],
        ["Dynamic", "Live web content and empirical evidence continuously change as new clinical trials and news emerge."],
        ["Multi-Agent", "Multiple specialized agents collaborate (Normalizer, Guardrail) and compete (PRO vs. CON) during factual verification."],
        ["Mixed Cooperative–Competitive", "PRO and CON compete adversarially to stress-test claims, while Cross-Examiner and Judge cooperate toward truth."],
        ["Discrete", "State transitions operate over discrete token sequences, debate rounds, numbered citations, and categorical verdicts."]
    ]

    t_shape4 = slide4.shapes.add_table(len(env_rows) + 1, len(env_headers), Inches(0.9), Inches(1.7), Inches(11.5), Inches(4.9))
    format_minimal_table(t_shape4.table, env_widths, env_headers, env_rows)

    # =========================================================================
    # SLIDE 5: Multi-Agent Architecture & System Flow
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_pure_slide_background(slide5)
    add_slide_header(slide5, "Multi-Agent Architecture & System Flow", "REVIEW 2: 3 MARKS")

    agent_headers = ["Agent", "Role", "Interaction & Coordination"]
    agent_widths = [2.2, 3.8, 5.5]
    agent_rows = [
        ["Claim Normalizer", "Breaks complex claim into atomic sub-claims", "Sends normalized sub-claims to PRO and CON debaters"],
        ["PRO Debater", "Discovers supporting empirical evidence via search", "Challenges CON arguments; registers supporting citations [Ex]"],
        ["CON Debater", "Discovers refuting empirical counter-evidence via search", "Challenges PRO arguments; registers refuting citations [Ex]"],
        ["Cross-Examiner", "Audits reasoning fallacies & dialectical synthesis", "Extracts points of consensus vs. contention for Supreme Judge"],
        ["Citation Guardrail", "Validates 100% of cited evidence IDs", "Checks debate text tags against evidence registry keys; flags hallucinated IDs"],
        ["Supreme Judge", "Makes final evidence-grounded calibrated decision", "Produces categorical verdict + confidence score + uncertainty note"]
    ]

    t_shape5 = slide5.shapes.add_table(len(agent_rows) + 1, len(agent_headers), Inches(0.9), Inches(1.7), Inches(11.5), Inches(4.0))
    format_minimal_table(t_shape5.table, agent_widths, agent_headers, agent_rows)

    # System Flow Box (Minimal Text)
    f_box = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(6.0), Inches(11.5), Inches(0.8))
    f_box.fill.solid()
    f_box.fill.fore_color.rgb = COLOR_HEADER_BG
    f_box.line.color.rgb = COLOR_BORDER
    f_box.line.width = Pt(1.0)

    tb_f = slide5.shapes.add_textbox(Inches(1.0), Inches(6.05), Inches(11.3), Inches(0.7))
    tf_f = tb_f.text_frame
    p_ft = tf_f.paragraphs[0]
    p_ft.text = "System Flow:  Claim  ➔  Normalize  ➔  [PRO ⟷ CON Debate Loop]  ➔  Cross-Examine  ➔  Verify Citations  ➔  Judge  ➔  Calibrated Verdict"
    p_ft.font.name = "Calibri"
    p_ft.font.size = Pt(11.5)
    p_ft.font.bold = True
    p_ft.font.color.rgb = COLOR_TEXT_PRIMARY

    # =========================================================================
    # SLIDE 6: Algorithmic Modeling & Search Strategy:
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_pure_slide_background(slide6)
    add_slide_header(slide6, "Algorithmic Modeling & Search Strategy:", "REVIEW 1: 3 MARKS")

    algo_headers = ["Component", "Algorithm / Technique", "Purpose & Guarantee"]
    algo_widths = [2.3, 3.4, 5.8]
    algo_rows = [
        ["Claim Decomposition", "Atomic Sub-Claim Normalization", "Breaks complex claims into 1–3 independent claims to prevent multi-hop retrieval errors."],
        ["Query Generation", "Stance-Conditioned Reformulation", "Generates separate polarized PRO and CON search queries to ensure balanced evidence retrieval."],
        ["Information Retrieval", "Hybrid Search (BM25 + FAISS Dense + RRF)", "Combines sparse lexical matching and dense vector similarity using Reciprocal Rank Fusion (c=60)."],
        ["Source Credibility", "Epistemic Authority Scorer", "Calculates domain reliability weights (Tier-1 Journals, Government .gov, Tier-2 Wire News vs. Blogs)."],
        ["Adversarial Reasoning", "Cyclic StateGraph (LangGraph)", "Allows PRO ⟷ CON debate across multiple rounds with state persistence."],
        ["Citation Verification", "Evidence-ID Validation", "Checks [E1], [E2], etc. against the evidence pool, ensuring zero hallucinated citations."],
        ["Final Decision", "Calibrated Epistemic Judge", "Produces TRUE, FALSE, or UNVERIFIABLE with hedged confidence for genuinely ambiguous claims."]
    ]

    t_shape6 = slide6.shapes.add_table(len(algo_rows) + 1, len(algo_headers), Inches(0.9), Inches(1.7), Inches(11.5), Inches(5.0))
    format_minimal_table(t_shape6.table, algo_widths, algo_headers, algo_rows)

    # =========================================================================
    # SLIDE 7: Tool/Package:
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_pure_slide_background(slide7)
    add_slide_header(slide7, "Tool/Package:", "REVIEW 2: 3 MARKS")

    tool_headers = ["Tool / Package", "Purpose in the Project"]
    tool_widths = [3.2, 8.3]
    tool_rows = [
        ["Python 3.11", "Core programming language for the multi-agent system, typing, and benchmarking."],
        ["LangGraph", "Builds the cyclic multi-agent debate and cross-examination workflow using StateGraph."],
        ["Groq LPU Engine", "Provides fast LLM inference for the PRO/CON debater agents and the Supreme Judge model."],
        ["Tavily & DDGS", "Performs live web search with multi-tier fallback resilience for factual empirical evidence."],
        ["FAISS & scikit-learn", "Vector-based similarity search and cosine matrix indexing in the hybrid retrieval module."],
        ["Rank-BM25", "Computes sparse lexical retrieval scores for exact medical and scientific keyword matching."],
        ["Streamlit", "Provides the interactive multi-tab web interface for claim verification and benchmark display."],
        ["python-dotenv", "Loads API keys and runtime configuration from .env securely."],
        ["Pydantic", "Structured schema validation of agent state data and JSON outputs."],
        ["pytest", "Unit and graph-level testing suite verifying claim decomposition, guardrails, and graph execution."]
    ]

    t_shape7 = slide7.shapes.add_table(len(tool_rows) + 1, len(tool_headers), Inches(0.9), Inches(1.7), Inches(11.5), Inches(5.0))
    format_minimal_table(t_shape7.table, tool_widths, tool_headers, tool_rows)

    # =========================================================================
    # SLIDE 8: Multi-Agent Execution & Interaction: (Clean Architectural Structure)
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_pure_slide_background(slide8)
    add_slide_header(slide8, "Multi-Agent Execution & Interaction:", "REVIEW 2: 3 MARKS")

    # 6 Clean Process Cards (Minimal borders, no loud colors)
    steps = [
        ("1. Input & Normalization", "Raw claim is decomposed into atomic sub-claims by Claim Normalizer.", Inches(0.9), Inches(1.8)),
        ("2. Adversarial Debate", "PRO and CON debaters iteratively retrieve live web evidence in rounds.", Inches(4.8), Inches(1.8)),
        ("3. Dialectical Audit", "Cross-Examiner checks reasoning fallacies and isolates empirical contention.", Inches(8.7), Inches(1.8)),
        ("4. Citation Guardrail", "Deterministic regex verifies [Ex] tags against the active evidence registry.", Inches(0.9), Inches(3.4)),
        ("5. Epistemic Arbitration", "Supreme Judge weighs source authority and evidence balance.", Inches(4.8), Inches(3.4)),
        ("6. Calibrated Verdict", "Final categorical verdict + confidence score + explicit uncertainty note.", Inches(8.7), Inches(3.4)),
    ]

    for title, desc, left, top in steps:
        box = slide8.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(3.7), Inches(1.3))
        box.fill.solid()
        box.fill.fore_color.rgb = COLOR_BG_LIGHT
        box.line.color.rgb = COLOR_BORDER
        box.line.width = Pt(1.0)

        tb = slide8.shapes.add_textbox(left + Inches(0.15), top + Inches(0.1), Inches(3.4), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.name = "Calibri"
        p_t.font.size = Pt(12)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_TEXT_PRIMARY

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.name = "Calibri"
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = COLOR_TEXT_SECONDARY

    # Understated Architectural Highlights Box
    tb_b = slide8.shapes.add_textbox(Inches(0.9), Inches(5.0), Inches(11.5), Inches(1.8))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True

    p_bt = tf_b.paragraphs[0]
    p_bt.text = "Key Architectural Benefits & Implementation Highlights:"
    p_bt.font.name = "Calibri"
    p_bt.font.size = Pt(12)
    p_bt.font.bold = True
    p_bt.font.color.rgb = COLOR_TEXT_PRIMARY

    h_bullets = [
        "Modular Multi-Agent Architecture: StateGraph allows cyclic debate turns with state persistence across rounds without context loss.",
        "Dialectical Triad: Thesis (PRO) + Antithesis (CON) + Synthesis (Cross-Examiner) prevents confirmation bias common in single-agent LLMs.",
        "Deterministic Citation Guardrail: Guarantees zero hallucinated evidence by rejecting ungrounded citation tags prior to judicial arbitration."
    ]

    for b in h_bullets:
        p = tf_b.add_paragraph()
        p.text = f"•  {b}"
        p.font.name = "Calibri"
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT_SECONDARY
        p.space_after = Pt(4)

    # =========================================================================
    # SLIDE 9: Output: (Testing Scenarios & Calibration Matrix)
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_pure_slide_background(slide9)
    add_slide_header(slide9, "Output: Demo Quality & Testing Scenarios", "REVIEW 2: 3 MARKS")

    eval_headers = ["Scenario", "Category", "Claim Statement", "Expected", "Output Verdict", "Confidence", "Guardrail Status"]
    eval_widths = [1.0, 2.0, 4.0, 1.2, 1.3, 1.0, 1.0]
    eval_rows = [
        ["1", "Clear False", "The Great Wall of China is visible from space with the naked eye.", "FALSE", "FALSE", "95%", "Clean (100%)"],
        ["2", "Clear True", "Regular exercise reduces the risk of cardiovascular disease.", "TRUE", "TRUE", "92%", "Clean (100%)"],
        ["3", "Ambiguous", "Moderate coffee consumption increases the risk of heart disease.", "UNVERIFIABLE", "UNVERIFIABLE", "55%", "Clean (Hedged)"],
        ["4", "Misleading", "Global average temperature increase of 1.5C has no impact on weather.", "FALSE", "FALSE", "94%", "Clean (100%)"],
        ["5", "Time-Sensitive", "Global lithium ion battery production volume doubled in 24 months.", "UNVERIFIABLE", "UNVERIFIABLE", "62%", "Clean (Hedged)"]
    ]

    t_shape9 = slide9.shapes.add_table(len(eval_rows) + 1, len(eval_headers), Inches(0.9), Inches(1.7), Inches(11.5), Inches(3.6))
    format_minimal_table(t_shape9.table, eval_widths, eval_headers, eval_rows)

    tb_notes = slide9.shapes.add_textbox(Inches(0.9), Inches(5.6), Inches(11.5), Inches(1.4))
    tf_notes = tb_notes.text_frame
    tf_notes.word_wrap = True

    p_nt = tf_notes.paragraphs[0]
    p_nt.text = "Empirical Validation & Calibration Notes:"
    p_nt.font.name = "Calibri"
    p_nt.font.size = Pt(11.5)
    p_nt.font.bold = True
    p_nt.font.color.rgb = COLOR_TEXT_PRIMARY

    p_n1 = tf_notes.add_paragraph()
    p_n1.text = "• 100% Benchmark Pass Rate: All 5 canonical test categories correctly reproduce ground-truth expectations across diverse domains."
    p_n1.font.name = "Calibri"
    p_n1.font.size = Pt(10.5)
    p_n1.font.color.rgb = COLOR_TEXT_SECONDARY

    p_n2 = tf_notes.add_paragraph()
    p_n2.text = "• Epistemic Nuance Calibration: Genuinely ambiguous or contested claims (e.g., coffee dosage, market statistics) do not collapse into binary True/False; confidence is calibrated to 55–62% with an explicit Uncertainty Advisory Note."
    p_n2.font.name = "Calibri"
    p_n2.font.size = Pt(10.5)
    p_n2.font.color.rgb = COLOR_TEXT_SECONDARY

    # Save Minimal Deck
    prs.save(output_filename)
    print(f"[Success] Minimal academic presentation generated: {output_filename}")


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "Aletheia_Dialectical_Fact_Checking.pptx"
    create_minimal_deck(out)
